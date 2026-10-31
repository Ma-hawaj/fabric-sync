"""Workbook-reading regressions; run with unittest discover -s scripts."""

import datetime as dt
import os
import tempfile
import unittest
import uuid
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import psycopg2
from psycopg2 import sql

from ingest_xlsm import (
    Report, measurement_signature, parse_measurement, perform_insert, read_workbook,
)


class InvoiceInsertTests(unittest.TestCase):
    def test_invalid_placeholder_quantity_is_rejected_before_database_access(self):
        for value in ("0", "-1", "NaN", "Infinity", "0.001", "100000000"):
            with self.subTest(value=value):
                conn = MagicMock()
                with self.assertRaises(ValueError):
                    perform_insert(conn, {}, "branch-id", Report(),
                                   material_amount=Decimal(value))
                conn.cursor.assert_not_called()

    def test_import_saves_ten_percent_vat_without_changing_charged_total(self):
        for gross, net, vat in [
            ("28.00", "25.45", "2.55"),
            ("11.00", "10.00", "1.00"),
            ("0.06", "0.05", "0.01"),
            ("0.00", "0.00", "0.00"),
        ]:
            with self.subTest(gross=gross):
                conn = MagicMock()
                cur = conn.cursor.return_value.__enter__.return_value
                cur.fetchone.side_effect = [("material-id",), ("invoice-id",)]
                cur.fetchall.return_value = [("Customer", "+97333333333", "customer-id")]
                workbook = {
                    "measurements": [], "repairs": [], "meas_to_invoices": {},
                    "invoices": [{
                        "invoice_number": 13481, "name": "Customer",
                        "phone": "+97333333333", "final": Decimal(gross),
                        "discount": Decimal("2.00"), "delivery": None,
                        "invoice_date": dt.date(2026, 1, 8),
                        "advance": Decimal("0"), "status": "pending", "meas_nos": [],
                    }],
                }
                state = {
                    "customer_rows": [("Customer", "+97333333333", "customer-id")],
                    "measurement_sigs": set(), "stages": {},
                    "orders_by_invoice": {}, "invoice_numbers": set(), "repairs": set(),
                }
                with patch("ingest_xlsm.load_db_state", return_value=state):
                    perform_insert(conn, workbook, "branch-id", Report())
                insert = next(call for call in cur.execute.call_args_list
                              if "INSERT INTO invoices" in call.args[0])
                sql, params = insert.args
                columns = sql.split("(", 1)[1].split(")", 1)[0]
                # These required tax columns have no database defaults.
                for column in ("vat_rate", "taxable_amount", "vat_amount"):
                    self.assertIn(column, columns)
                self.assertEqual(sql.count("%s"), len(params))
                self.assertEqual(params[2], Decimal(gross))
                self.assertEqual(params[6], Decimal("2.00"))
                self.assertEqual(params[-3:],
                                 (Decimal("0.10"), Decimal(net), Decimal(vat)))
                self.assertEqual(sum(params[-2:]), params[2])


class ParseMeasurementTests(unittest.TestCase):
    def test_symbols_attached_spaced_and_standalone(self):
        for symbol, fraction in [
            ("½", "0.5"), ("¼", "0.25"), ("¾", "0.75"),
            ("⅛", "0.125"), ("⅜", "0.375"), ("⅝", "0.625"), ("⅞", "0.875"),
        ]:
            for text, whole in [(symbol, 0), (f"55{symbol}", 55),
                                (f"55  {symbol}", 55), (f"55\u00a0{symbol}", 55)]:
                with self.subTest(text=text):
                    self.assertEqual(parse_measurement(text),
                                     (Decimal(whole) + Decimal(fraction), True))

    def test_slash_fractions(self):
        for text, expected in [("55 1/2", "55.5"), ("17 3/4", "17.75"),
                               ("⅛", "0.125"), ("1/8", "0.125"),
                               ("17 3 / 8", "17.375"), ("17 5⁄8", "17.625")]:
            with self.subTest(text=text):
                self.assertEqual(parse_measurement(text), (Decimal(expected), True))

    def test_existing_workbook_notation(self):
        for text, expected in [("15 ½ .1", "15.5"), ("1 ¼.1", "1.25"),
                               ("24. ½", "24.5"), ("18,5", "18.5"),
                               ("7. 75", "7.75"), ("24. ⅜", "24.375"),
                               ("17⅝.1", "17.625"), (32, "32"),
                               (17.125, "17.125"), (Decimal("17.875"), "17.875")]:
            with self.subTest(text=text):
                self.assertEqual(parse_measurement(text), (Decimal(expected), True))

    def test_blank_invalid_and_ambiguous_values(self):
        for text in [None, "", "   "]:
            self.assertEqual(parse_measurement(text), (None, True))
        for text in ["CANCEL", "SHIRT", "17-18", "17 ½ ¾", "1/0", "3/2",
                     "1/3", "55 1/2 extra", float("nan"), float("inf"),
                     Decimal("NaN"), True]:
            with self.subTest(text=text):
                self.assertEqual(parse_measurement(text), (None, False))

    def test_rerun_signatures_preserve_eighths_and_ignore_scale(self):
        date = dt.date(2026, 5, 1)
        def signature(value):
            return measurement_signature("customer", date, {"chest": value})
        self.assertEqual(signature(Decimal("42.125")), signature(Decimal("42.1250")))
        self.assertEqual(signature(42), signature(Decimal("42.000")))
        self.assertNotEqual(signature(Decimal("42.125")), signature(Decimal("42.13")))


class ReadWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.date = dt.datetime(2026, 5, 1)
        self.invoices = [
            [no, self.date, "Receipt Customer", "33333333", "7 & 8 & 9", 30]
            for no in (101, "102", 103.0)
        ]
        self.measurements = [
            [no, self.date, "Sheet Customer", "34444444", 42, "29", "3"]
            for no in (7, "8", 9.0)
        ]

    def read(self, invoices=None, measurements=None):
        wb = openpyxl.Workbook()
        wb.active.title = "New"
        wb["New"].append([
            "Invoice", "Invoice Date", "Name", "Telephone", "Measurement No",
            "Final Amount",
        ])
        for row in self.invoices if invoices is None else invoices:
            wb["New"].append(row)
        ws = wb.create_sheet("Measurments")
        ws.append([
            "Measurement No", "Date", "Customer Name", "Customer Telephone",
            "Chest", "Button", "Button Fold",
        ])
        for row in self.measurements if measurements is None else measurements:
            ws.append(row)
        wb.create_sheet("Repairing").append(["RECEIPT DATE", "Measurement No"])
        path = Path(self.tempdir.name) / "orders.xlsm"
        wb.save(path)
        wb.close()
        report = Report()
        return read_workbook(path, report), report

    def test_invalid_invoice_numbers_are_reported_and_valid_rows_unchanged(self):
        expected, _ = self.read()
        rows = list(self.invoices)
        for index, value in enumerate((None, "", "   ", "not-a-number", self.date)):
            rows.insert(index * 2, [
                value, self.date, "Invalid Receipt", "35555555", "99", 10,
            ])
        actual, report = self.read(invoices=rows)
        self.assertEqual(actual, expected)
        self.assertEqual([inv["invoice_number"] for inv in actual["invoices"]],
                         [101, 102, 103])
        # Empty Excel strings load as None, retaining the existing anomaly.
        self.assertEqual(report.numeric["New: dated row without invoice number"], 2)
        self.assertEqual(report.numeric["New: invalid invoice number skipped"], 3)
        self.assertIn("not-a-number", report.summary("workbook"))

    def test_invalid_measurement_numbers_are_reported_and_valid_rows_unchanged(self):
        expected, _ = self.read()
        rows = list(self.measurements)
        for index, value in enumerate((None, "", "   ", "not-a-number", self.date)):
            rows.insert(index * 2, [
                value, self.date, "Invalid Measurement", "35555555", 40,
            ])
        actual, report = self.read(measurements=rows)
        self.assertEqual(actual, expected)
        self.assertEqual([m["measurement_no"] for m in actual["measurements"]],
                         [7, 8, 9])
        self.assertEqual(report.numeric[
            "Measurments: invalid measurement number skipped"
        ], 5)
        self.assertIn("not-a-number", report.summary("workbook"))

    def test_measurement_duplicate_selection_is_preserved(self):
        rows = [
            [7, dt.datetime(2026, 1, 1), "Customer", "33333333", 40],
            [7, self.date, "Customer", "33333333", 42],
            [7, dt.datetime(2026, 4, 1), "Customer", "33333333", 41],
            [7, dt.datetime(2026, 6, 1), "Customer", "33333333", 43],
        ]
        actual, report = self.read(invoices=[], measurements=rows)
        self.assertEqual(len(actual["measurements"]), 1)
        measurement = actual["measurements"][0]
        self.assertEqual(measurement["measurement_date"], dt.date(2026, 6, 1))
        self.assertEqual(measurement["values"]["chest"], Decimal("43"))
        self.assertEqual(report.numeric["Measurments: edit kept over placeholder"], 2)
        self.assertEqual(report.numeric["Measurments: duplicate number dropped"], 1)

    def test_button_columns_are_imported_as_bottom(self):
        # The sheet misspells Bottom / Bottom Fold as Button / Button Fold;
        # the values are hem widths and land on the bottom columns.
        actual, _ = self.read()
        self.assertEqual(len(actual["measurements"]), 3)
        for measurement in actual["measurements"]:
            self.assertEqual(measurement["values"]["bottom"], Decimal("29"))
            self.assertEqual(
                measurement["values"]["bottom_fold"], Decimal("3"))
            self.assertNotIn("button", measurement["values"])
            self.assertNotIn("button_fold", measurement["values"])

    def test_fraction_cells_are_imported_without_losing_eighths(self):
        rows = [
            [7, self.date, "Customer", "33333333", "42 ½", "29¼", "3 ⅛"],
            [8, self.date, "Customer", "33333333", "42 3/8", "29¾", "3⅞"],
        ]
        actual, _ = self.read(measurements=rows)
        expected = [("42.5", "29.25", "3.125"), ("42.375", "29.75", "3.875")]
        for measurement, values in zip(actual["measurements"], expected, strict=True):
            for key, value in zip(["chest", "bottom", "bottom_fold"], values):
                self.assertEqual(measurement["values"][key], Decimal(value))


@unittest.skipUnless(os.environ.get("TEST_DATABASE_URL"),
                     "set TEST_DATABASE_URL to a disposable PostgreSQL 18+ database")
class DatabaseIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.conn = psycopg2.connect(os.environ["TEST_DATABASE_URL"])
        self.addCleanup(self.conn.close)
        self.addCleanup(self.conn.rollback)
        self.root = Path(__file__).resolve().parents[1]
        # All DDL/data remains in this transaction and is rolled back on cleanup.
        schema = "ingest_test_" + uuid.uuid4().hex
        with self.conn.cursor() as cur:
            cur.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))
            cur.execute(sql.SQL("SET LOCAL search_path TO {}, public")
                        .format(sql.Identifier(schema)))
            for migration in sorted((self.root / "backend/migrations").glob("*.sql")):
                cur.execute(migration.read_text())
            cur.execute((self.root / "backend/seeds/reference_data.sql").read_text())

    def test_workbook_insert_payments_repairs_and_idempotent_rerun(self):
        fixture = ReadWorkbookTests()
        fixture.setUp()
        self.addCleanup(fixture.tempdir.cleanup)
        workbook, report = fixture.read()
        invoice = workbook["invoices"][0]
        invoice.update(final=Decimal("28.00"), discount=Decimal("2.00"),
                       advance=Decimal("10.00"), status="done", pay1="cash")
        workbook["repairs"] = [{
            "meas_nos": [7], "reported_on": dt.date(2026, 6, 1),
            "reason": "Adjust sleeve", "charge": Decimal("1.00"),
            "raw_status": "done", "collected": dt.datetime(2026, 6, 2),
        }]
        with self.conn.cursor() as cur:
            cur.execute("INSERT INTO branch (name) VALUES ('Import test') RETURNING id")
            branch = cur.fetchone()[0]
        perform_insert(self.conn, workbook, branch, report,
                       material_amount=Decimal("2.50"))
        with self.conn.cursor() as cur:
            cur.execute("SELECT total_price, vat_rate, taxable_amount, vat_amount "
                        "FROM invoices WHERE invoice_number = 101")
            self.assertEqual(cur.fetchone(), tuple(map(Decimal,
                             ("28.00", "0.10", "25.45", "2.55"))))
            cur.execute("SELECT sum(amount) FROM invoice_payments")
            self.assertEqual(cur.fetchone()[0], Decimal("28.00"))
            cur.execute("SELECT count(*) FROM order_repairs")
            self.assertEqual(cur.fetchone()[0], 1)
            cur.execute("SELECT count(*), sum(price) FROM orders")
            self.assertEqual(cur.fetchone(), (9, Decimal("88.00")))
            cur.execute("SELECT min(material_amount), max(material_amount) FROM orders")
            self.assertEqual(cur.fetchone(), (Decimal("2.50"), Decimal("2.50")))
        perform_insert(self.conn, workbook, branch, Report())
        with self.conn.cursor() as cur:
            for table, count in (("invoices", 3), ("orders", 9),
                                 ("invoice_payments", 2), ("order_repairs", 1)):
                cur.execute(sql.SQL("SELECT count(*) FROM {}").format(sql.Identifier(table)))
                self.assertEqual(cur.fetchone()[0], count, table)

    def test_seed_sets_all_emd_prices_and_preserves_staff_changes(self):
        seed = (self.root / "backend/seeds/design_options.sql").read_text()
        with self.conn.cursor() as cur:
            cur.execute(seed)
            cur.execute("SELECT count(*), min(cost), max(cost) FROM design_options "
                        "WHERE category = 'emd'")
            self.assertEqual(cur.fetchone(), (27, Decimal("5.00"), Decimal("5.00")))
            cur.execute("SELECT count(*) FROM design_options WHERE category <> 'emd' AND cost <> 0")
            self.assertEqual(cur.fetchone()[0], 0)
            cur.execute("UPDATE design_options SET cost = 7 WHERE category = 'emd' AND slug = '1'")
            cur.execute(seed)
            cur.execute("SELECT count(*), count(*) FILTER (WHERE cost = 5), "
                        "count(*) FILTER (WHERE cost = 7) FROM design_options WHERE category = 'emd'")
            self.assertEqual(cur.fetchone(), (27, 26, 1))


if __name__ == "__main__":
    unittest.main()
